"""
Carga datos reales MinTIC 2024 e ISPs por ciudad desde archivos xlsx
Uso: python manage.py cargar_mintic_isps
"""
from django.core.management.base import BaseCommand
from django.db import transaction
from api.models import MunicipioData
import json


class Command(BaseCommand):
    help = 'Carga datos MinTIC 2024 e ISPs directorio desde xlsx'

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.stdout.write(self.style.ERROR('openpyxl no instalado'))
            return

        self.stdout.write('Procesando BASE_ISP.xlsx...')
        isps_por_ciudad = self.cargar_isps_ciudad()

        self.stdout.write('Procesando Datos_ISPs_Colombia.xlsx (puede tardar 1-2 min)...')
        accesos_por_dane = self.cargar_accesos_mintic()

        self.stdout.write('Actualizando municipios en BD...')
        self.actualizar_municipios(isps_por_ciudad, accesos_por_dane)

    def cargar_isps_ciudad(self):
        from openpyxl import load_workbook
        from collections import defaultdict

        wb = load_workbook('/opt/cobertura_isp/data/BASE ISP.xlsx', read_only=True)
        ws = wb['ISP']
        isps = defaultdict(list)

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] and row[5]:
                ciudad = str(row[5]).strip().upper()
                if 'BOGOT' in ciudad:
                    ciudad = 'BOGOTÁ'
                isps[ciudad].append({
                    'empresa': str(row[1]).strip()[:100],
                    'contacto': str(row[2]).strip()[:80] if row[2] else '',
                    'telefono': str(row[3]).strip()[:20] if row[3] else '',
                    'correo': str(row[4]).strip()[:80] if row[4] else '',
                })
        wb.close()
        self.stdout.write(f'  {len(isps)} ciudades, {sum(len(v) for v in isps.values())} ISPs')
        return dict(isps)

    def cargar_accesos_mintic(self):
        from openpyxl import load_workbook
        from collections import defaultdict

        wb = load_workbook('/opt/cobertura_isp/data/Datos_ISPs_Colombia.xlsx', read_only=True)
        ws = wb['4,3']

        accesos = defaultdict(lambda: {
            'accesos': 0, 'proveedores': set(),
            'departamento': '', 'municipio': ''
        })

        header_found = False
        count = 0
        for row in ws.iter_rows(values_only=True):
            if not header_found:
                if row[0] == 'AÑO':
                    header_found = True
                continue
            if row[5] and row[11]:
                try:
                    cod = str(row[5]).strip().zfill(5)
                    accesos[cod]['accesos'] += int(row[11] or 0)
                    if row[2]:
                        accesos[cod]['proveedores'].add(str(row[2]).strip())
                    accesos[cod]['departamento'] = str(row[4]).strip() if row[4] else ''
                    accesos[cod]['municipio'] = str(row[6]).strip() if row[6] else ''
                    count += 1
                except:
                    pass

        wb.close()
        resultado = {
            k: {
                'accesos': v['accesos'],
                'proveedores_count': len(v['proveedores']),
                'proveedores': list(v['proveedores'])[:20],
                'departamento': v['departamento'],
                'municipio': v['municipio'],
            }
            for k, v in accesos.items()
        }
        self.stdout.write(f'  {count:,} filas procesadas, {len(resultado)} municipios')
        return resultado

    def actualizar_municipios(self, isps_ciudad, accesos_dane):
        actualizados = 0
        sin_mintic = 0

        with transaction.atomic():
            for m in MunicipioData.objects.all():
                nombre = m.nombre.upper().strip()

                # 1. Datos MinTIC por código DANE
                mintic = accesos_dane.get(m.codigo_dane)
                if mintic:
                    m.mintic_accesos_reales = mintic['accesos']
                    m.mintic_proveedores_reales = mintic['proveedores_count']
                    m.mintic_proveedores_count = mintic['proveedores_count']
                    m.mintic_tiene_fibra = mintic['proveedores_count'] > 0
                    if m.hogares and m.hogares > 0 and mintic['accesos'] > 0:
                        m.mintic_penetracion_pct = round(
                            min(99.9, mintic['accesos'] / m.hogares * 100), 1
                        )
                else:
                    sin_mintic += 1

                # 2. ISPs directorio por nombre de ciudad
                isps = isps_ciudad.get(nombre, [])
                if not isps:
                    # Buscar coincidencia parcial
                    for key in isps_ciudad:
                        if nombre[:5] in key or key[:5] in nombre:
                            isps = isps_ciudad[key]
                            break

                m.isps_directorio = json.dumps(isps[:15], ensure_ascii=False)
                m.save()
                actualizados += 1

        self.stdout.write(self.style.SUCCESS(
            f'\n✅ {actualizados} municipios actualizados'
            f'\n   Sin datos MinTIC: {sin_mintic}'
        ))

        # Verificar ejemplos
        for cod, nombre in [('05001', 'Medellín'), ('11001', 'Bogotá'), ('47001', 'Santa Marta')]:
            try:
                m = MunicipioData.objects.get(codigo_dane=cod)
                isps_count = len(json.loads(m.isps_directorio or '[]'))
                self.stdout.write(
                    f'   {nombre}: {m.mintic_accesos_reales or 0:,} accesos, '
                    f'{m.mintic_proveedores_count} ISPs MinTIC, '
                    f'{isps_count} ISPs directorio, '
                    f'{m.mintic_penetracion_pct}% penetración'
                )
            except:
                pass
